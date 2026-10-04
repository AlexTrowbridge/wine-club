"""Build small workbooks in the legacy 'Wine Club.xlsx' layout for tests."""

import datetime as dt

from openpyxl import Workbook


def build_workbook(path, members, eras):
    """eras: [(era_name, [session, ...])], session = dict(number, date, theme, bottles,
    attendance={member: "P"/"X"}), bottle = dict(name, purchaser, scores={member: n|"x"},
    position, price, blend, region, country, jh)."""
    wb = Workbook()
    raw = wb.active
    raw.title = "Raw Ratings"
    res = wb.create_sheet("Results")
    att = wb.create_sheet("Attendance")

    raw.append([None, None, None, None, "Ratings"])
    raw.append(["#", "Wine", "Date", "Purchased by", "Average", *members])
    res.append(["Wine Details", None, None, None, None, "Ratings"])
    res.append(
        [
            "#",
            "Wine",
            "Date",
            "Wine #",
            "Purchased by",
            "Average",
            "Age Consumed",
            "Price",
            "Single Varietal/Blend",
            "Region",
            "Country",
            "Min",
            "Max",
            "MinMax Diff",
            "100 pt scale",
            "JH Score",
        ]
    )
    att.append([])
    att.append([None, None, None, *members])
    att.append([None, None, None, *[0] * len(members)])
    att.append(["No.", "Date", "Theme", *[0] * len(members)])

    n = 0
    for era_name, sessions in eras:
        raw.append([None, era_name])
        res.append([None, era_name])
        att.append([era_name])
        for s in sessions:
            date = dt.datetime.combine(s["date"], dt.time())
            att.append(
                [
                    s["number"],
                    date,
                    s.get("theme", "-"),
                    *[s.get("attendance", {}).get(m) for m in members],
                ]
            )
            for b in s["bottles"]:
                n += 1
                scores = [b["scores"].get(m) for m in members]
                nums = [v for v in scores if isinstance(v, (int, float))]
                avg = sum(nums) / len(nums) if nums else None
                raw.append([n, b["name"], date, b["purchaser"], avg, *scores])
                res.append(
                    [
                        n,
                        b["name"],
                        date,
                        b.get("position", 1),
                        b["purchaser"],
                        avg,
                        b.get("age"),
                        b.get("price"),
                        b.get("blend", "Single Varietal"),
                        b.get("region"),
                        b.get("country"),
                        min(nums, default=None),
                        max(nums, default=None),
                        None,
                        None,
                        b.get("jh"),
                    ]
                )
    # Trailing pre-numbered empty rows, as in the real workbook.
    for k in range(n + 1, n + 4):
        raw.append([k])
        res.append([k, None, None, None, None, " "])
    wb.save(path)


MEMBERS = ["Dion", "Jun", "Alex", "Zane"]

SAMPLE = [
    (
        "Old Era",
        [
            {
                "number": 1,
                "date": dt.date(2023, 4, 21),
                "attendance": {"Dion": "P", "Jun": "X", "Alex": "X"},
                "bottles": [
                    {
                        "name": "Elysian Springs 2017 Spring Lamb Syrah",
                        "purchaser": "Dion",
                        "scores": {"Dion": 9, "Jun": 7, "Alex": 9},
                        "price": 20,
                        "region": "Adelaide Hills",
                        "country": "Australia",
                        "age": 6,
                    },
                ],
            },
        ],
    ),
    (
        "New Era",
        [
            {
                "number": 2,
                "date": dt.date(2023, 5, 26),
                "theme": "Grenache Day",
                "attendance": {"Dion": "X", "Jun": "P", "Alex": "P", "Zane": "X"},
                "bottles": [
                    {
                        "name": "Henschke 2020 Henry's Seven Shiraz Grenache Mataro Viognier",
                        "purchaser": "Jun",
                        "scores": {"Dion": 8, "Jun": 8, "Alex": 9, "Zane": "x"},
                        "position": 1,
                        "price": 33,
                        "blend": "Blend",
                        "region": "Barossa Valley",
                        "country": "Australia",
                        "jh": 92,
                    },
                    {
                        "name": "Bollinger NV Special Cuvée Champagne Pinot Noir Chardonnay",
                        "purchaser": "Everyone",
                        "scores": {"Dion": 9, "Jun": 9, "Alex": 8},
                        "position": 2,
                        "price": 90,
                        "blend": "Blend",
                        "region": "Champagne",
                        "country": "France",
                    },
                    {
                        "name": "Pepperjack 2021 Cabernet Sauvignon",
                        "purchaser": "Alex",
                        "scores": {"Dion": 6, "Jun": 7, "Alex": 7, "Zane": 6},
                        "position": 3,
                        "price": 18,
                        "region": "Barossa Valley",
                        "country": "Australia",
                    },
                ],
            },
        ],
    ),
]
